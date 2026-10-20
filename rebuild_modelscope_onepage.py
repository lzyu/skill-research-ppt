"""One-page ModelScope: 3 capabilities on top, 4 open ecosystem blocks below."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from redesign_lifecycle import add_rect, add_text, set_run

SRC = "skill-final-v1.pptx"
RED = RGBColor(0xB5, 0x33, 0x3B)
TITLE_RED = RGBColor(0xC7, 0x00, 0x0B)
INK = RGBColor(0x1A, 0x1A, 0x1C)
GRAY = RGBColor(0x5B, 0x61, 0x6B)
LIGHT_GRAY = RGBColor(0x8B, 0x8E, 0x94)
BORDER = RGBColor(0xE2, 0xE4, 0xE7)
TRACK = RGBColor(0xED, 0xEF, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x56, 0x64, 0x72)
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def add_list(s, x, y, w, h, items, size=9.5):
    box = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(3)
        r1 = p.add_run(); r1.text = "· "
        set_run(r1, size, RED, True)
        r2 = p.add_run(); r2.text = text
        set_run(r2, size, INK, False)
    return box


def add_bar(s, x, y, w, h, ratio, color):
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, TRACK, None, radius=0.4)
    fill_w = max(int(w * ratio), int(0.002 * w)) if ratio > 0 else 0
    if fill_w:
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, fill_w, h, color, None, radius=0.4)


def build(s, W, H):
    add_text(s, 0.035 * W, 0.040 * H, 0.93 * W, 0.055 * H,
             "魔搭 Skill 关键能力：发布两步、安装四路，试运行和原创供给都还薄",
             24, TITLE_RED, True)

    # Top: three capability cards
    caps = [
        ("01", "发布", "两步 API，约束不少",
         ["传 zip → 建 Skill；包 ≤5MB，只能 1 个 SKILL.md",
          "name / owner 不可改；审核、version 不透明"]),
        ("02", "使用", "四条安装通道，搜索不好用",
         ["npx / SDK / bash / zip；默认 ~/.agents/skills",
          "OpenAPI 关键词 search 不生效，只能按类目 / 作者筛"]),
        ("03", "试运行", "没有技能级沙箱",
         ["只有「前往实验场试用」→ 通用 Agent 会话",
          "无输入 / 结果 / 回放；MCP 有实验场，Skill 没有"]),
    ]
    for i, (num, title, sub, items) in enumerate(caps):
        x = (0.035 + i * 0.320) * W
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.112 * H, 0.300 * W, 0.180 * H, WHITE, BORDER, 0.75, radius=0.035)
        add_text(s, x + 0.014 * W, 0.124 * H, 0.030 * W, 0.022 * H, num, 9, RED, True)
        add_text(s, x + 0.042 * W, 0.122 * H, 0.10 * W, 0.024 * H, title, 11, INK, True)
        add_text(s, x + 0.014 * W, 0.150 * H, 0.27 * W, 0.020 * H, sub, 8.5, GRAY, False)
        add_list(s, x + 0.014 * W, 0.175 * H, 0.27 * W, 0.100 * H, items, 9.5)
    add_text(s, 0.675 * W + 0.014 * W, 0.266 * H, 0.27 * W, 0.020 * H,
             "实测：iic/fun-denoise 5.3k 浏览 / 782 下载", 8, LIGHT_GRAY, False)

    # Ecosystem: open 2x2 comparison, no boxes
    add_text(s, 0.035 * W, 0.312 * H, 0.40 * W, 0.026 * H, "生态画像：4 项对比", 12, INK, True)

    # 01 supply
    add_text(s, 0.035 * W, 0.345 * H, 0.46 * W, 0.026 * H, "供给：数量很大，原创很薄", 11, INK, True)
    supply = [
        ("平台官方", 337, "0.4%", 337 / 83509, RED),
        ("品牌 / 垂直", 550, "<1%", 550 / 83509, SLATE),
        ("聚合 / 镜像", 25808, "30.9%", 25808 / 83509, RGBColor(0x9A, 0xA1, 0xAA)),
        ("社区长尾", 56814, "68.0%", 56814 / 83509, RGBColor(0xCF, 0xD4, 0xDA)),
    ]
    for i, (name, count, share, ratio, color) in enumerate(supply):
        y = 0.385 * H + i * 0.032 * H
        add_text(s, 0.035 * W, y, 0.100 * W, 0.020 * H, name, 9.5, GRAY, False)
        add_bar(s, 0.140 * W, y + 0.004 * H, 0.160 * W, 0.011 * H, ratio, color)
        add_text(s, 0.310 * W, y, 0.165 * W, 0.020 * H, f"{count:,} / {share}", 9, INK, False)
    add_text(s, 0.035 * W, 0.515 * H, 0.46 * W, 0.034 * H,
             "官方 337 条；垂直约 550；镜像 25,808 条；其余约 5.68 万（按总数估算）。", 8.5, LIGHT_GRAY, False, spacing=1.05)

    # 02 consumption
    add_text(s, 0.520 * W, 0.345 * H, 0.44 * W, 0.026 * H, "消费：上架即沉默", 11, INK, True)
    kpis = [("158", "中位下载（样本 2,400）"), ("65.9%", "Top1% 技能拿走下载"), ("3.5×", "品牌/垂直中位 vs 长尾")]
    for i, (num, label) in enumerate(kpis):
        x = 0.520 * W + i * 0.145 * W
        add_text(s, x, 0.385 * H, 0.130 * W, 0.040 * H, num, 18, RED, True)
        add_text(s, x, 0.425 * H, 0.130 * W, 0.034 * H, label, 8.5, GRAY, False, spacing=1.05)
    add_text(s, 0.520 * W, 0.500 * H, 0.44 * W, 0.040 * H,
             "下载含 CLI / 镜像拉取，不等于用户数；样本偏向推荐位，真实中位更低。", 8.5, LIGHT_GRAY, False, spacing=1.05)

    # 03 category
    add_text(s, 0.035 * W, 0.605 * H, 0.46 * W, 0.026 * H, "类目：Top3 占 66.3%", 11, INK, True)
    add_text(s, 0.035 * W, 0.630 * H, 0.46 * W, 0.020 * H,
             "10 个一级类目；35 个二级类目里 24 个是 0 条", 8.5, LIGHT_GRAY, False)
    cats = [("开发工具", 27928), ("代码质检", 16854), ("媒体处理", 10569),
            ("前端开发", 6686), ("云效工具", 5544)]
    for i, (name, count) in enumerate(cats):
        y = 0.655 * H + i * 0.026 * H
        add_text(s, 0.035 * W, y, 0.100 * W, 0.020 * H, name, 9, GRAY, False)
        add_bar(s, 0.140 * W, y + 0.004 * H, 0.190 * W, 0.010 * H, count / 27928,
                RED if i == 0 else RGBColor(0xC9, 0x7A, 0x80))
        add_text(s, 0.345 * W, y, 0.080 * W, 0.020 * H, f"{count:,}", 8.5, INK, False)
    add_text(s, 0.035 * W, 0.790 * H, 0.46 * W, 0.020 * H,
             "开发工具 + 代码质检 + 媒体处理 = 55,351 条", 8.5, LIGHT_GRAY, False)

    # 04 who gets downloaded
    add_text(s, 0.520 * W, 0.605 * H, 0.46 * W, 0.026 * H,
             "谁被下载：品牌 > 官方 > 长尾 > 镜像", 11, INK, True)
    roles = [("品牌 / 垂直", 516), ("平台官方", 374), ("社区长尾", 148), ("聚合 / 镜像", 145)]
    for i, (name, value) in enumerate(roles):
        y = 0.645 * H + i * 0.032 * H
        add_text(s, 0.520 * W, y, 0.100 * W, 0.020 * H, name, 9, GRAY, False)
        add_bar(s, 0.640 * W, y + 0.004 * H, 0.150 * W, 0.011 * H, value / 516,
                RED if i == 0 else RGBColor(0x9A, 0xA1, 0xAA))
        add_text(s, 0.805 * W, y, 0.060 * W, 0.020 * H, str(value), 8.5, INK, False)
    add_text(s, 0.520 * W, 0.775 * H, 0.46 * W, 0.020 * H,
             "最高异常：1.15M 下载 / 592 浏览；下载不宜直接当用户数。", 8.5, LIGHT_GRAY, False)


def main():
    prs = Presentation(SRC)
    W, H = prs.slide_width, prs.slide_height
    old_id = prs.slides[1].slide_id

    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.shapes):
        ph._element.getparent().remove(ph._element)
    build(s, W, H)
    new_sld = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(new_sld)
    prs.slides._sldIdLst.insert(1, new_sld)

    old_sld = next(x for x in list(prs.slides._sldIdLst) if x.id == old_id)
    prs.part.drop_rel(old_sld.get(R_NS + "id"))
    prs.slides._sldIdLst.remove(old_sld)

    prs.save(SRC)
    print("saved", SRC, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
