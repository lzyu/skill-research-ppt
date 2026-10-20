"""Replace the single ModelScope page with two pages: publish/use/trial + ecosystem."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from redesign_lifecycle import add_rect, add_text, set_run

SRC = "skill-final-v1.pptx"
RED = RGBColor(0xB5, 0x33, 0x3B)
TITLE_RED = RGBColor(0xC7, 0x00, 0x0B)
INK = RGBColor(0x1A, 0x1A, 0x1C)
GRAY = RGBColor(0x5B, 0x61, 0x6B)
LIGHT_GRAY = RGBColor(0x8B, 0x8E, 0x94)
BORDER = RGBColor(0xE2, 0xE4, 0xE7)
BAND = RGBColor(0xF7, 0xF8, 0xFA)
TRACK = RGBColor(0xED, 0xEF, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x56, 0x64, 0x72)
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def add_list(s, x, y, w, h, items, size=10):
    box = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(4)
        r1 = p.add_run(); r1.text = "· "
        set_run(r1, size, RED, True)
        r2 = p.add_run(); r2.text = text
        set_run(r2, size, INK, False)
    return box


def add_bar(s, x, y, w, h, ratio, color):
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, TRACK, None, radius=0.4)
    fill_w = max(int(w * ratio), int(0.002 * w)) if ratio > 0 else 0
    if fill_w:
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, fill_w, h, color, None, radius=0.4)


def add_page_a(prs, W, H):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.shapes):
        ph._element.getparent().remove(ph._element)

    add_text(s, 0.04 * W, 0.045 * H, 0.92 * W, 0.06 * H,
             "魔搭 Skill 中心关键能力：两步发布、四路安装，试运行还停在通用实验场",
             24, TITLE_RED, True)

    cols = [
        ("01", "发布", "两步 API，约束不少",
         ["① 上传 zip → ② 创建 Skill",
          "根目录只能有 1 个 SKILL.md，包 ≤5MB",
          "name / owner 创建后不可改",
          "GitHub 导入或本地 zip；审核状态不公开",
          "没有独立 version 字段，版本落在 git 分支"]),
        ("02", "使用", "四条安装通道，搜索却不好用",
         ["npx / SDK / bash / 手动 zip 四条通道",
          "默认装到 ~/.agents/skills",
          "Cursor / Claude Code / Codex / Gemini 可发现",
          "OpenAPI 关键词 search 不生效，只能按类目 / 作者筛",
          "私有 Skill 安装要 --token，且 token 按站点隔离"]),
        ("03", "试运行", "没有技能级沙箱",
         ["详情页只有「前往实验场试用」",
          "跳到 /agents/playground，通用 Agent 会话",
          "没有示例输入、执行结果、回放",
          "替代：本地装完跑 / 读 examples / MS-Agent test_skill",
          "对比：MCP 有实验场，Skill 没有"]),
    ]
    xs = [0.035, 0.355, 0.675]
    for i, (num, name, sub, items) in enumerate(cols):
        x = xs[i] * W
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.13 * H, 0.29 * W, 0.38 * H, WHITE, BORDER, 0.75, radius=0.04)
        add_text(s, x + 0.014 * W, 0.145 * H, 0.030 * W, 0.028 * H, num, 12, RED, True)
        add_text(s, x + 0.044 * W, 0.145 * H, 0.10 * W, 0.028 * H, name, 12, INK, True)
        add_text(s, x + 0.014 * W, 0.178 * H, 0.26 * W, 0.025 * H, sub, 10, GRAY, False)
        add_list(s, x + 0.014 * W, 0.212 * H, 0.26 * W, 0.27 * H, items, 10)

    breaks = [
        ("发布", "审核不透明，版本很弱", "无审核状态字段；无 version 落库；source_url 是自由文本。"),
        ("使用", "搜索只到类目 / 作者", "OpenAPI search 未生效；第 25 页 403；站内搜索是内部接口。"),
        ("试运行", "没有技能级执行入口", "实验场是通用 Agent 会话；无输入 / 结果 / 回放。"),
    ]
    for i, (tag, head, body) in enumerate(breaks):
        x = xs[i] * W
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.545 * H, 0.29 * W, 0.155 * H, BAND, BORDER, 0.75, radius=0.04)
        add_text(s, x + 0.014 * W, 0.558 * H, 0.035 * W, 0.02 * H, tag, 9, RED, True)
        add_text(s, x + 0.055 * W, 0.558 * H, 0.22 * W, 0.025 * H, head, 12, INK, True)
        add_text(s, x + 0.014 * W, 0.590 * H, 0.262 * W, 0.09 * H, body, 10, GRAY, False, spacing=1.1)

    # Screenshot evidence band: detail page + install section.
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.035 * W, 0.735 * H, 0.93 * W, 0.170 * H, BAND, BORDER, 0.75, radius=0.035)
    hero = s.shapes.add_picture("modelscope-detail-hero.png", int(0.045 * W), int(0.745 * H), width=int(0.38 * W))
    hero.line.color.rgb = BORDER
    hero.line.width = Pt(0.75)
    install = s.shapes.add_picture("modelscope-install.png", int(0.445 * W), int(0.745 * H), width=int(0.140 * W))
    install.line.color.rgb = BORDER
    install.line.width = Pt(0.75)

    add_text(s, 0.610 * W, 0.747 * H, 0.345 * W, 0.028 * H,
             "详情页：5.3k 浏览 / 782 下载，主按钮只到 /agents/playground。", 9, INK, False, spacing=1.15)
    add_text(s, 0.610 * W, 0.785 * H, 0.345 * W, 0.028 * H,
             "安装区：npx / SDK / bash / zip 四条命令，只有复制，没有一键上云。", 9, GRAY, False, spacing=1.15)
    add_text(s, 0.610 * W, 0.823 * H, 0.345 * W, 0.050 * H,
             "详情页没有“运行”面板；“Skill 文件”只能浏览和逐文件下载；发布与试跑都要登录。", 9, GRAY, False, spacing=1.15)

    return s


def add_page_b(prs, W, H):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.shapes):
        ph._element.getparent().remove(ph._element)

    add_text(s, 0.04 * W, 0.045 * H, 0.92 * W, 0.06 * H,
             "魔搭 Skill 生态画像：8.35 万个 Skill，官方 0.4%，头部 1% 拿走 66% 下载",
             24, TITLE_RED, True)

    # Supply structure
    add_text(s, 0.04 * W, 0.125 * H, 0.54 * W, 0.030 * H, "供给：数量很大，原创很薄", 12, INK, True)
    supply = [
        ("平台官方", 337, "0.4%", 337 / 83509, RED),
        ("品牌 / 垂直", "~550", "<1%", 550 / 83509, SLATE),
        ("聚合 / 镜像", 25808, "30.9%", 25808 / 83509, RGBColor(0x9A, 0xA1, 0xAA)),
        ("社区长尾", 56814, "68.0%", 56814 / 83509, RGBColor(0xCF, 0xD4, 0xDA)),
    ]
    for i, (name, count, share, ratio, color) in enumerate(supply):
        y = 0.165 * H + i * 0.038 * H
        add_text(s, 0.04 * W, y, 0.10 * W, 0.022 * H, name, 10, GRAY, False)
        add_bar(s, 0.145 * W, y + 0.004 * H, 0.255 * W, 0.014 * H, ratio, color)
        add_text(s, 0.410 * W, y, 0.165 * W, 0.022 * H, (f"{count:,} / {share}" if isinstance(count, int) else f"{count} / {share}"), 9.5, INK, False)
    add_text(s, 0.04 * W, 0.322 * H, 0.54 * W, 0.024 * H,
             "官方 337 条；垂直约 550；镜像 25,808 条；其余约 5.68 万（按总数估算）", 9, LIGHT_GRAY, False)

    # Category bars
    add_text(s, 0.04 * W, 0.365 * H, 0.54 * W, 0.028 * H, "类目：Top3 占 66.3%", 12, INK, True)
    add_text(s, 0.04 * W, 0.392 * H, 0.54 * W, 0.022 * H,
             "10 个一级类目；35 个二级类目里 24 个是 0 条", 9, LIGHT_GRAY, False)
    cats = [
        ("开发工具", 27928), ("代码质检", 16854), ("媒体处理", 10569),
        ("前端开发", 6686), ("云效工具", 5544),
    ]
    for i, (name, count) in enumerate(cats):
        y = 0.425 * H + i * 0.040 * H
        add_text(s, 0.04 * W, y, 0.10 * W, 0.022 * H, name, 10, GRAY, False)
        add_bar(s, 0.145 * W, y + 0.004 * H, 0.255 * W, 0.014 * H, count / 27928, RED if i == 0 else RGBColor(0xC9, 0x7A, 0x80))
        add_text(s, 0.410 * W, y, 0.10 * W, 0.022 * H, f"{count:,}", 9.5, INK, False)
    add_text(s, 0.04 * W, 0.630 * H, 0.54 * W, 0.022 * H,
             "开发工具 + 代码质检 + 媒体处理 = 55,351 条", 9, LIGHT_GRAY, False)

    # Consumption KPIs
    add_text(s, 0.61 * W, 0.125 * H, 0.35 * W, 0.030 * H, "消费：上架即沉默", 12, INK, True)
    kpis = [
        ("158", "中位下载（样本 2,400）"),
        ("65.9%", "Top1% 技能拿走下载"),
        ("3.5×", "品牌/垂直中位 vs 长尾"),
    ]
    for i, (num, label) in enumerate(kpis):
        x = 0.61 * W + i * 0.115 * W
        add_text(s, x, 0.160 * H, 0.105 * W, 0.045 * H, num, 20, RED, True)
        add_text(s, x, 0.207 * H, 0.105 * W, 0.055 * H, label, 9, GRAY, False, spacing=1.1)
    add_text(s, 0.61 * W, 0.275 * H, 0.35 * W, 0.040 * H,
             "下载含 CLI / 镜像拉取，不等于用户数；样本偏向推荐位，真实中位更低。", 9, LIGHT_GRAY, False)

    # Median downloads by role
    add_text(s, 0.61 * W, 0.365 * H, 0.35 * W, 0.028 * H, "谁被下载：品牌 > 官方 > 长尾 > 镜像", 12, INK, True)
    roles = [
        ("品牌 / 垂直", 516), ("平台官方", 374), ("社区长尾", 148), ("聚合 / 镜像", 145),
    ]
    for i, (name, value) in enumerate(roles):
        y = 0.405 * H + i * 0.045 * H
        add_text(s, 0.61 * W, y, 0.09 * W, 0.022 * H, name, 10, GRAY, False)
        add_bar(s, 0.705 * W, y + 0.004 * H, 0.16 * W, 0.014 * H, value / 516, RED if i == 0 else RGBColor(0x9A, 0xA1, 0xAA))
        add_text(s, 0.875 * W, y, 0.08 * W, 0.022 * H, str(value), 9.5, INK, False)
    add_text(s, 0.61 * W, 0.590 * H, 0.35 * W, 0.040 * H,
             "最高异常：1.15M 下载 / 592 浏览；下载不宜直接当用户数。", 9, LIGHT_GRAY, False)

    # Mirror explanation strip
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.04 * W, 0.735 * H, 0.92 * W, 0.160 * H, BAND, BORDER, 0.75, radius=0.04)
    add_text(s, 0.055 * W, 0.747 * H, 0.22 * W, 0.024 * H, "30.9% 的“镜像”是什么", 10, RED, True)
    add_text(s, 0.055 * W, 0.775 * H, 0.88 * W, 0.050 * H,
             "4 个批量搬运账号：majiayu000 16,093、diegosouzapw 9,315、beatra01 211、smyx 189；合计 25,808 条。不是逐条去重，也不是官方口径；74.3% 的 source_url 指向 GitHub。",
             9.5, GRAY, False, spacing=1.2)
    add_text(s, 0.055 * W, 0.830 * H, 0.88 * W, 0.050 * H,
             "镜像不等于重复；数量涨了，消费没涨（聚合中位 145）；官方 337 条是信任锚点，品牌/垂直中位 516。",
             9.5, GRAY, False, spacing=1.2)

    return s


def main():
    prs = Presentation(SRC)
    W, H = prs.slide_width, prs.slide_height
    # Keep the old page until the two new slides have been created. Adding after
    # deletion can reuse a part name and produce duplicate slide20.xml entries.
    old_id = prs.slides[1].slide_id

    add_page_a(prs, W, H)
    a_sld = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(a_sld)
    prs.slides._sldIdLst.insert(1, a_sld)

    add_page_b(prs, W, H)
    b_sld = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(b_sld)
    prs.slides._sldIdLst.insert(2, b_sld)

    old_sld = next(x for x in list(prs.slides._sldIdLst) if x.id == old_id)
    prs.part.drop_rel(old_sld.get(R_NS + "id"))
    prs.slides._sldIdLst.remove(old_sld)

    prs.save(SRC)
    print("saved", SRC, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
