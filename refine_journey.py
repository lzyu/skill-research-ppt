"""Rebuild page 1: concrete actions + checks + outputs for the Skill lifecycle."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from redesign_lifecycle import add_rect, add_text, set_run

SRC = "skill-final-v1.pptx"
RED = RGBColor(0xB5, 0x33, 0x3B)
TITLE_RED = RGBColor(0xC7, 0x00, 0x0B)
INK = RGBColor(0x1A, 0x1A, 0x1C)
GRAY = RGBColor(0x5B, 0x61, 0x6B)
LIGHT_GRAY = RGBColor(0x8B, 0x8E, 0x94)
BORDER = RGBColor(0xE2, 0xE4, 0xE7)
BAND = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_actions(s, x, y, w, h, actions):
    box = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, action in enumerate(actions):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.0
        p.space_after = Pt(3)
        r1 = p.add_run(); r1.text = "· "
        set_run(r1, 10, RED, True)
        r2 = p.add_run(); r2.text = action
        set_run(r2, 10, INK, False)
    return box


def main():
    prs = Presentation(SRC)
    W, H = prs.slide_width, prs.slide_height
    s = prs.slides[0]
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)

    add_text(s, 0.04 * W, 0.045 * H, 0.92 * W, 0.06 * H,
             "Skill 资产生命周期：生产 → 分发 → 安装 → 使用",
             24, TITLE_RED, True)

    rail_y = 0.160 * H
    rail = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(0.06 * W), int(rail_y), int(0.94 * W), int(rail_y))
    rail.line.color.rgb = RGBColor(0xC9, 0xCD, 0xD2)
    rail.line.width = Pt(1.0)

    stages = [
        ("01", "生产", "作者", "把经验变成可复用的方法",
         ["定任务：目标 / 输入输出 / 边界",
          "起稿：$skill-creator 生成草稿",
          "补步骤/异常/验收，配资源跑三类任务"],
         "触发准 / 边界清 / 样例与评测齐全",
         "方法包 + 样例 + 评测 + 依赖"),
        ("02", "分发", "作者 / 平台", "让合适的人拿到明确版本",
         ["打包：Git / ZIP，附来源/版本/许可",
          "写说明页：场景/示例/环境/限制",
          "上架：市场 / Plugin / 托管 API"],
         "展示=实得 / 来源可追 / 版本可回滚",
         "发布包 + 说明页 + 来源/版本/许可"),
        ("03", "安装", "宿主 Agent", "在宿主 Agent 中安装",
         ["复制到 ~/.agents/skills",
          "npx skills add <repo> -g -a codex",
          "查依赖 / 权限，跑最小任务"],
         "宿主可见 / 依赖权限 / 最小任务通过",
         "安装实例 + 就绪检查"),
        ("04", "使用", "用户 / Agent", "在任务中调用并反馈",
         ["提任务：目标 / 输入 / 约束 / 验收",
          "自动匹配，或点名 /pdf-processing",
          "渐进加载 references/scripts"],
         "已加载 / 已执行 / 已验收",
         "任务结果 + 验收结论 + 使用记录"),
    ]

    card_w, gap = 0.22 * W, 0.0167 * W
    card_x0, card_y, card_h = 0.035 * W, 0.200 * H, 0.350 * H
    centers = []
    for i, (num, name, role, verb, actions, check, product) in enumerate(stages):
        x = card_x0 + i * (card_w + gap)
        centers.append(x + card_w / 2)
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_w, card_h, WHITE, BORDER, 0.75, radius=0.05)
        add_text(s, x + 0.014 * W, card_y + 0.010 * H, 0.030 * W, 0.030 * H, num, 12, RED, True)
        add_text(s, x + 0.040 * W, card_y + 0.010 * H, 0.085 * W, 0.030 * H, name, 12, INK, True)
        add_text(s, x + 0.125 * W, card_y + 0.012 * H, 0.080 * W, 0.025 * H, role, 8.5, LIGHT_GRAY,
                 False, PP_ALIGN.RIGHT)
        add_text(s, x + 0.014 * W, card_y + 0.044 * H, 0.190 * W, 0.032 * H, verb, 12, INK, True)

        add_text(s, x + 0.014 * W, card_y + 0.078 * H, 0.040 * W, 0.018 * H, "动作", 8.5, RED, True)
        add_actions(s, x + 0.014 * W, card_y + 0.096 * H, 0.190 * W, 0.125 * H, actions)

        # divider + check + product
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x + 0.012 * W), int(card_y + 0.228 * H),
                                      int(x + card_w - 0.012 * W), int(card_y + 0.228 * H))
        line.line.color.rgb = BORDER
        line.line.width = Pt(0.5)

        add_text(s, x + 0.014 * W, card_y + 0.238 * H, 0.035 * W, 0.018 * H, "检查", 8.5, RED, True)
        add_text(s, x + 0.050 * W, card_y + 0.238 * H, 0.160 * W, 0.018 * H, check, 9, GRAY, False)
        add_text(s, x + 0.014 * W, card_y + 0.272 * H, 0.035 * W, 0.018 * H, "产物", 8.5, RED, True)
        add_text(s, x + 0.050 * W, card_y + 0.272 * H, 0.160 * W, 0.026 * H, product, 9, GRAY, False)

    # Nodes and arrows
    node_d = 0.028 * H
    for i, cx in enumerate(centers):
        node = add_rect(s, MSO_SHAPE.OVAL, cx - node_d / 2, rail_y - node_d / 2, node_d, node_d, WHITE, RED, 2.0)
        tf = node.text_frame
        tf.word_wrap = False
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = stages[i][0]
        set_run(r, 9, RED, True)
        if i < 3:
            ax = cx + node_d / 2 + 0.018 * W
            add_rect(s, MSO_SHAPE.RIGHT_ARROW, ax, rail_y - 0.007 * H, 0.017 * W, 0.014 * H, RED, None, radius=0.3)

    # Feedback loop
    loop_y = 0.580 * H
    loop = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(0.93 * W), int(loop_y), int(0.07 * W), int(loop_y))
    loop.line.color.rgb = RGBColor(0xA5, 0xAC, 0xB4)
    loop.line.width = Pt(1)
    loop.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_rect(s, MSO_SHAPE.LEFT_ARROW, 0.05 * W, loop_y - 0.007 * H, 0.017 * W, 0.014 * H, RGBColor(0xA5, 0xAC, 0xB4), None, radius=0.3)
    label = add_text(s, 0.36 * W, loop_y - 0.016 * H, 0.28 * W, 0.026 * H,
                     "用后反馈 → 修订 → 新版本", 9.5, GRAY, False, PP_ALIGN.CENTER)
    label.fill.solid(); label.fill.fore_color.rgb = WHITE; label.line.fill.background()

    # Zhiguoyuan example strip
    strip_x, strip_y, strip_w, strip_h = 0.035 * W, 0.615 * H, 0.93 * W, 0.200 * H
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, strip_x, strip_y, strip_w, strip_h, BAND, BORDER, 0.75, radius=0.04)
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, strip_x + 0.008 * W, strip_y + 0.018 * H,
             0.004 * W, strip_h - 0.036 * H, RED, None, radius=0.5)
    add_text(s, strip_x + 0.020 * W, strip_y + 0.010 * H, 0.20 * W, 0.022 * H,
             "智果园示例", 9, RED, True)
    add_text(s, strip_x + 0.020 * W, strip_y + 0.030 * H, 0.55 * W, 0.028 * H,
             "从生产到 CodeArts Agent 的一条链路", 12, INK, True)

    steps = [
        ("01", "生产", "本地或 CodeArts 生产 Skill"),
        ("02", "上架", "智果园 AI 资产广场"),
        ("03", "安装", "CodeArts Agent"),
        ("04", "调用", "任务中按需加载"),
        ("05", "反馈", "使用记录 → v1.1"),
    ]
    step_w, step_gap = 0.155 * W, 0.022 * W
    step_x0 = strip_x + 0.020 * W
    for i, (num, label, body) in enumerate(steps):
        x = step_x0 + i * (step_w + step_gap)
        add_text(s, x, strip_y + 0.064 * H, 0.026 * W, 0.020 * H, num, 8.5, RED, True)
        add_text(s, x + 0.026 * W, strip_y + 0.064 * H, 0.045 * W, 0.020 * H, label, 9.5, INK, True)
        add_text(s, x, strip_y + 0.086 * H, step_w, 0.040 * H, body, 9.5, GRAY, False, spacing=1.0)
        if i < len(steps) - 1:
            add_rect(s, MSO_SHAPE.RIGHT_ARROW, x + step_w + 0.004 * W, strip_y + 0.086 * H,
                     0.010 * W, 0.010 * H, RED, None, radius=0.3)

    add_text(s, strip_x + 0.020 * W, strip_y + 0.158 * H, 0.88 * W, 0.024 * H,
             "验收：能装 · 能跑 · 能验收 · 能回滚", 9, GRAY, False)

    # Closing judgment
    add_text(s, 0.04 * W, 0.855 * H, 0.045 * W, 0.022 * H, "判断", 9.5, RED, True)
    add_text(s, 0.088 * W, 0.855 * H, 0.87 * W, 0.024 * H,
             "平台的价值不在“存资产”，而在让 Skill 可发现、可安装、可运行、可迭代。",
             10, GRAY, False)

    prs.save(SRC)
    print("saved", SRC, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
