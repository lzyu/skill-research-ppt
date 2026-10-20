"""Merge selected Skill research/commercialization decks into one PPTX.

The source decks are all 16:9 (12192000 x 6858000). Exact duplicate files are
represented once; distinct versions (v1/v2, alternative designs) are kept.
No source files are modified.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from pptx import Presentation

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
BASE = "modelscope-skills-center.pptx"
OUT = "skill-final-v1.pptx"

# section, source file, expected slides
SOURCES = [
    ("01 魔搭平台调研", "modelscope-skills-center.pptx", 1),
    ("02 SkillHub 平台调研", "skillhub/output/skillhub-insights.pptx", 1),
    ("02 SkillHub 平台调研", "skillhub/output/skillhub-insights-v1.pptx", 1),
    ("02 SkillHub 平台调研", "skillhub-eval-onepage.pptx", 1),
    ("03 Skill 平台总结", "skill-asset-lifecycle.pptx", 1),
    ("03 Skill 平台总结", "zhiguoyuan-skill-asset-onepage.pptx", 1),
    ("04 腾讯 SkillPay 商业化", "skillhub/output/skillhub-skillpay-commercialization-v1.pptx", 2),
    ("04 腾讯 SkillPay 商业化", "skillhub/output/skillhub-skillpay-commercialization-v2-final.pptx", 2),
    ("05 阿里 SkillPay 商业化", "gpt-alipay-research.pptx", 4),
    ("05 阿里 SkillPay 商业化", "alipay-skillpay-machinepay.pptx", 4),
    ("06 Skill 商业化总结", "zhiguoyuan-skill-commercialization.pptx", 2),
]


def clear_slides(prs):
    for sld_id in list(prs.slides._sldIdLst):
        r_id = sld_id.get(R_NS + "id")
        prs.part.drop_rel(r_id)
        prs.slides._sldIdLst.remove(sld_id)


def copy_slide(src_slide, dst_prs):
    # The Huawei template has a single layout named Blank; remove its placeholders.
    dst = dst_prs.slides.add_slide(dst_prs.slide_layouts[0])
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)

    for shp in src_slide.shapes:
        el = deepcopy(shp._element)
        # Re-map image relationships to the destination package.
        for blip in el.xpath(".//a:blip"):
            r_id = blip.get(R_NS + "embed")
            if r_id:
                rel = src_slide.part.rels[r_id]
                if rel.is_external:
                    continue
                part = rel.target_part
                _, new_r_id = dst.part.get_or_add_image_part(BytesIO(part.blob))
                blip.set(R_NS + "embed", new_r_id)
        dst.shapes._spTree.insert_element_before(el, "p:extLst")

    # Notes are intentionally not copied: notes slides inherit a theme-specific
    # notes master, and cross-theme notes relationships can make the merged deck
    # unstable in some renderers. The source decks keep their original notes.


def main():
    base = Presentation(BASE)
    width, height = base.slide_width, base.slide_height
    clear_slides(base)

    merged = []
    for section, path, expected in SOURCES:
        src = Presentation(path)
        assert (src.slide_width, src.slide_height) == (width, height), f"size mismatch: {path}"
        assert len(src.slides) == expected, f"slide count mismatch: {path}: {len(src.slides)} != {expected}"
        for slide in src.slides:
            copy_slide(slide, base)
            merged.append((section, path, len(merged) + 1))

    base.core_properties.title = "Skill 资产与商业化研究合集 v1"
    base.core_properties.subject = "魔搭 / SkillHub / SkillPay / 支付宝 / 智果园"
    base.save(OUT)

    check = Presentation(OUT)
    assert len(check.slides) == len(merged), f"merged slide count {len(check.slides)} != {len(merged)}"
    print(f"saved {OUT}: {len(check.slides)} slides, {Path(OUT).stat().st_size // 1024} KB")
    for i, (section, path, _) in enumerate(merged, 1):
        s = check.slides[i - 1]
        pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
        texts = sum(1 for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        print(f"{i:02d} {section:24s} {Path(path).name:55s} shapes={len(s.shapes):3d} pics={pics:2d} texts={texts:3d}")


if __name__ == "__main__":
    main()
