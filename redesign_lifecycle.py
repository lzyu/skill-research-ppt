"""Remove title-area microcopy and rebuild the Skill lifecycle page as page 1."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC = "skill-final-v1.pptx"
FONT = "微软雅黑"
RED = RGBColor(0xB5, 0x33, 0x3B)
TITLE_RED = RGBColor(0xC7, 0x00, 0x0B)
INK = RGBColor(0x1A, 0x1A, 0x1C)
GRAY = RGBColor(0x5B, 0x61, 0x6B)
LIGHT_GRAY = RGBColor(0x8B, 0x8E, 0x94)
BORDER = RGBColor(0xE2, 0xE4, 0xE7)
BAND = RGBColor(0xF7, 0xF8, 0xFA)
DARK = RGBColor(0x1F, 0x29, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def set_run(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def add_text(slide, x, y, w, h, text, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    box = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    set_run(r, size, color, bold)
    return box


def add_bullets(slide, x, y, w, h, bullets):
    box = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.08
        p.space_after = Pt(3)
        r1 = p.add_run()
        r1.text = "· "
        set_run(r1, 10, RED, True)
        r2 = p.add_run()
        r2.text = text
        set_run(r2, 10, INK, False)
    return box


def add_rect(slide, shape_type, x, y, w, h, fill, line=None, line_w=0.75, radius=None):
    shp = slide.shapes.add_shape(shape_type, int(x), int(y), int(w), int(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    if radius is not None and shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def max_size(shape):
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else 0


def text_of(shape):
    return shape.text_frame.text.strip() if shape.has_text_frame else ""


def remove(shape):
    shape._element.getparent().remove(shape._element)


def clean_title_area(slide, W, H):
    for shape in list(slide.shapes):
        if not shape.has_text_frame:
            # red rules are line shapes
            if shape.shape_type == 9:  # LINE
                try:
                    rgb = shape.line.color.rgb
                except Exception:
                    rgb = None
                if rgb is not None and str(rgb).upper() == "B5333B":
                    top = shape.top / H
                    if 0.10 <= top <= 0.18 and shape.width / W >= 0.5:
                        remove(shape)
            continue
        text = text_of(shape)
        if not text:
            continue
        size = max_size(shape)
        left, top = shape.left / W, shape.top / H
        # Kicker above the title.
        if top < 0.06 and left < 0.20 and shape.width / W >= 0.40 and size <= 12:
            remove(shape)
            continue
        # One-line summary directly under the title.
        if 0.115 <= top <= 0.17 and left <= 0.06 and shape.width / W >= 0.35 and size <= 11.5 and len(text) >= 15:
            remove(shape)
            continue
        # Red rule drawn as a very thin rectangle.
        if 0.10 <= top <= 0.18 and shape.width / W >= 0.5 and shape.height / H <= 0.012:
            try:
                rgb = shape.fill.fore_color.rgb
            except Exception:
                rgb = None
            if rgb is not None and str(rgb).upper() in ("B5333B", "C7000B"):
                remove(shape)


def build_lifecycle(prs, layout, W, H):
    s = prs.slides.add_slide(layout)
    for ph in list(s.shapes):
        ph._element.getparent().remove(ph._element)

    # Title
    add_text(s, 0.04 * W, 0.045 * H, 0.92 * W, 0.06 * H,
             "Skill 的四段旅程：写出来 → 发出去 → 装进去 → 用起来",
             24, TITLE_RED, True)

    # Journey band and rail
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.025 * W, 0.215 * H, 0.95 * W, 0.375 * H, BAND, None, radius=0.04)
    rail_y = 0.19 * H
    rail = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(0.06 * W), int(rail_y), int(0.94 * W), int(rail_y))
    rail.line.color.rgb = RED
    rail.line.width = Pt(1.25)

    stages = [
        ("01", "生产", "作者", "把方法写下来",
         ["SKILL.md 写清 name/description、步骤、异常与验收",
          "用 $skill-creator 生成草稿，再拿真实任务改",
          "留一个能跑通的样例，写清“不该用”的边界"],
         "产物：SKILL.md + scripts/ + references/"),
        ("02", "分发", "作者 / 平台", "把版本交出去",
         ["Git 仓库保留目录与提交历史，ref 指向实际 commit",
          "ZIP/目录包离线交付，附依赖与完整性说明",
          "市场目录 / Plugin 组合交付，标明版本与许可"],
         "产物：发布包 + 来源 / 版本 / 许可"),
        ("03", "安装", "宿主", "让宿主发现",
         ["npx skills add <repo> --skill <name> -g -a codex",
          "复制到 ~/.agents/skills 或 ~/.claude/skills",
          "查依赖与权限，跑一个最小任务验证"],
         "产物：安装实例 + 就绪检查"),
        ("04", "使用", "用户 / Agent", "把结果交付",
         ["提目标、输入、约束和验收标准",
          "渐进加载：目录 → SKILL.md → references/ → scripts/",
          "记录版本/环境/修正；先分清方法、输入还是环境"],
         "产物：任务结果 + 使用记录"),
    ]

    card_w, gap = 0.22 * W, 0.0167 * W
    card_x0, card_y, card_h = 0.035 * W, 0.24 * H, 0.34 * H
    centers = []
    for i, (num, name, role, verb, bullets, product) in enumerate(stages):
        x = card_x0 + i * (card_w + gap)
        centers.append(x + card_w / 2)
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_w, card_h, WHITE, BORDER, 0.75, radius=0.05)
        add_text(s, x + 0.014 * W, card_y + 0.012 * H, 0.03 * W, 0.03 * H, num, 12, RED, True)
        add_text(s, x + 0.040 * W, card_y + 0.012 * H, 0.09 * W, 0.03 * H, name, 12, INK, True)
        add_text(s, x + 0.135 * W, card_y + 0.014 * H, 0.07 * W, 0.025 * H, role, 9, LIGHT_GRAY,
                 False, PP_ALIGN.RIGHT)
        add_text(s, x + 0.014 * W, card_y + 0.052 * H, 0.19 * W, 0.035 * H, verb, 12, INK, True)
        add_bullets(s, x + 0.014 * W, card_y + 0.094 * H, 0.19 * W, 0.175 * H, bullets)
        add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.012 * W, card_y + 0.292 * H,
                 0.196 * W, 0.038 * H, RGBColor(0xF5, 0xF6, 0xF8), None, radius=0.12)
        add_text(s, x + 0.020 * W, card_y + 0.300 * H, 0.18 * W, 0.028 * H, product, 9, GRAY, False,
                 spacing=1.0)

    # Nodes and arrows
    node_d = 0.032 * H
    for i, cx in enumerate(centers):
        node = add_rect(s, MSO_SHAPE.OVAL, cx - node_d / 2, rail_y - node_d / 2, node_d, node_d, WHITE, RED, 2.0)
        tf = node.text_frame
        tf.word_wrap = False
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = stages[i][0]
        set_run(r, 10, RED, True)
        if i < 3:
            ax = cx + node_d / 2 + 0.018 * W
            add_rect(s, MSO_SHAPE.RIGHT_ARROW, ax, rail_y - 0.008 * H, 0.018 * W, 0.016 * H, RED, None, radius=0.3)

    # Return loop
    loop_y = 0.615 * H
    loop = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(0.93 * W), int(loop_y), int(0.07 * W), int(loop_y))
    loop.line.color.rgb = RGBColor(0xA5, 0xAC, 0xB4)
    loop.line.width = Pt(1)
    loop.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_rect(s, MSO_SHAPE.LEFT_ARROW, 0.05 * W, loop_y - 0.008 * H, 0.018 * W, 0.016 * H, RGBColor(0xA5, 0xAC, 0xB4), None, radius=0.3)
    label = add_text(s, 0.36 * W, loop_y - 0.018 * H, 0.28 * W, 0.03 * H,
                     "用完复盘 → 修订 → 新版本", 10, GRAY, False, PP_ALIGN.CENTER)
    label.fill.solid(); label.fill.fore_color.rgb = WHITE; label.line.fill.background()

    # Case strip
    case_x, case_y, case_w, case_h = 0.035 * W, 0.66 * H, 0.93 * W, 0.27 * H
    add_rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, case_x, case_y, case_w, case_h, DARK, None, radius=0.03)
    add_text(s, case_x + 0.018 * W, case_y + 0.018 * H, 0.45 * W, 0.03 * H,
             "以 PDF 处理 Skill 为例", 12, WHITE, True)
    add_text(s, case_x + 0.70 * W, case_y + 0.020 * H, 0.20 * W, 0.025 * H,
             "一次完整旅程", 9, RGBColor(0x9C, 0xA3, 0xAF), False, PP_ALIGN.RIGHT)

    steps = [
        ("生产", "SKILL.md + scripts/"),
        ("分发", "打 v1.0 · 推 Git"),
        ("安装", "npx skills add"),
        ("使用", "/pdf-processing 合并 PDF"),
        ("验收", "页数 · 文本完整"),
        ("反馈", "补 OCR → v1.1"),
    ]
    step_w = 0.135 * W
    step_x0 = case_x + 0.010 * W
    step_gap = 0.018 * W
    for i, (label, body) in enumerate(steps):
        x = step_x0 + i * (step_w + step_gap)
        add_text(s, x, case_y + 0.060 * H, step_w, 0.022 * H, label, 9, RGBColor(0xE8, 0x7A, 0x82), True)
        add_text(s, x, case_y + 0.085 * H, step_w, 0.050 * H, body, 10, WHITE, False, spacing=1.0)
        if i < len(steps) - 1:
            add_rect(s, MSO_SHAPE.RIGHT_ARROW, x + step_w + 0.002 * W, case_y + 0.093 * H,
                     0.010 * W, 0.012 * H, RED, None, radius=0.3)

    # Bottom checkpoints inside the dark case strip.
    div_y = case_y + 0.196 * H
    div = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(case_x + 0.018 * W), int(div_y), int(case_x + case_w - 0.018 * W), int(div_y))
    div.line.color.rgb = RGBColor(0x3A, 0x46, 0x52)
    div.line.width = Pt(0.75)
    checks = [
        ("能装", "npx skills add"),
        ("能跑", "最小任务通过"),
        ("能验收", "页数 / 文本 / 版本"),
        ("能回滚", "v1.0 → v1.1"),
    ]
    cw = 0.20 * W
    for i, (k, v) in enumerate(checks):
        cx = case_x + 0.018 * W + i * 0.222 * W
        add_text(s, cx, div_y + 0.018 * H, cw, 0.02 * H, k, 9, RGBColor(0xE8, 0x7A, 0x82), True)
        add_text(s, cx, div_y + 0.040 * H, cw, 0.03 * H, v, 10, WHITE, False)

    return s


def main():
    prs = Presentation(SRC)
    W, H = prs.slide_width, prs.slide_height
    layout = prs.slide_layouts[0]

    old_lifecycle_id = prs.slides[4].slide_id

    # Clean every existing page's title area.
    for slide in prs.slides:
        clean_title_area(slide, W, H)

    # Build the new journey page at the end, then move it to the front.
    build_lifecycle(prs, layout, W, H)
    new_sld = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(new_sld)
    prs.slides._sldIdLst.insert(0, new_sld)

    # Remove the old lifecycle page by its stable slide id.
    old_sld = next(sld for sld in list(prs.slides._sldIdLst) if sld.id == old_lifecycle_id)
    prs.part.drop_rel(old_sld.get(R_NS + "id"))
    prs.slides._sldIdLst.remove(old_sld)

    # After removing the kicker, move each title up to the top margin.
    for slide in prs.slides:
        title = None
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            size = max_size(shape)
            if size >= 18 and shape.top / H < 0.13 and shape.left / W < 0.25:
                if title is None or size > title[0]:
                    title = (size, shape)
        if title:
            shape = title[1]
            shape.top = int(0.045 * H)
            shape.width = W - shape.left - int(0.04 * W)

    prs.core_properties.title = "Skill 资产与商业化研究合集 v1（旅程版）"
    prs.save(SRC)

    check = Presentation(SRC)
    print(f"saved {SRC}: {len(check.slides)} slides")
    for i, slide in enumerate(list(check.slides)[:4], 1):
        texts = [sh.text_frame.text.replace("\n", " ")[:40] for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        print(i, texts[:3])


if __name__ == "__main__":
    main()
