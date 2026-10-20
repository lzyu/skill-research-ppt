"""Remove thin rasterized rules that sit immediately under slide titles."""
from pptx import Presentation

SRC = "skill-final-v1.pptx"
p = Presentation(SRC)
W, H = p.slide_width, p.slide_height
removed = []
for i, s in enumerate(p.slides, 1):
    for sh in list(s.shapes):
        if sh.shape_type != 13:  # PICTURE
            continue
        top, width, height = sh.top / H, sh.width / W, sh.height / H
        if 0.08 <= top <= 0.20 and width >= 0.40 and height <= 0.12:
            removed.append((i, round(top, 3), round(width, 3), round(height, 3)))
            sh._element.getparent().remove(sh._element)
p.save(SRC)
print("removed", len(removed), "top rule pictures")
for r in removed:
    print(r)
